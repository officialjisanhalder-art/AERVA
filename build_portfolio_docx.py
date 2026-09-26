import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout_box(doc, text_list, title=None, border_color="0052FF", bg_color="F8FAFC"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:left w:val="single" w:sz="36" w:space="0" w:color="{border_color}"/>\n'
        f'  <w:top w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'  <w:bottom w:val="none"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    
    if title:
        run_title = p.add_run(f"★ {title}\n")
        run_title.bold = True
        run_title.font.name = 'Calibri'
        run_title.font.size = Pt(11)
        run_title.font.color.rgb = RGBColor(10, 37, 64)
    
    for idx, t in enumerate(text_list):
        if idx > 0 or title:
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
        run = p.add_run(t)
        run.font.name = 'Calibri'
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(30, 41, 59)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def style_table_header(row, bg_color="0A2540"):
    for cell in row.cells:
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.name = 'Calibri'
                run.font.size = Pt(10.5)

def build_portfolio():
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
    
    styles = doc.styles
    normal_style = styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Calibri'
    normal_font.size = Pt(11)
    normal_font.color.rgb = RGBColor(30, 41, 59)
    
    NAVY = RGBColor(10, 37, 64)
    BLUE = RGBColor(0, 82, 255)
    DARK = RGBColor(15, 23, 42)
    GRAY = RGBColor(100, 116, 139)
    
    # Header Section
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(2)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("EXECUTIVE PORTFOLIO & PROJECT SHOWCASE")
    r_title.bold = True
    r_title.font.size = Pt(24)
    r_title.font.color.rgb = NAVY
    
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(12)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Full-Stack Web & Mobile Solutions Architect | Product Systems Specialist")
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = BLUE
    r_sub.bold = True
    
    add_callout_box(doc, [
        "Specializing in high-impact mobile applications (Flutter/iOS/Android), custom web portals, real-time dispatch systems, and scalable cloud databases for startups, client projects, and corporate enterprises.",
        "Featured Award: 🥈 2nd Prize Winner — NearER Emergency Healthcare Platform."
    ], title="FULL-STACK DEVELOPER & SYSTEMS ARCHITECT FOR HIRE", border_color="0052FF", bg_color="F8FAFC")
    
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(15)
        r.font.color.rgb = NAVY
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(12)
        r.font.color.rgb = BLUE
        return p

    def add_body(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(text)
        r.font.size = Pt(10.5)
        r.font.color.rgb = DARK
        return p

    def add_bullet(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        r_b = p.add_run(bold_prefix + " ")
        r_b.bold = True
        r_b.font.color.rgb = NAVY
        r_t = p.add_run(text)
        r_t.font.color.rgb = DARK

    def add_image_visual(img_filename, caption_text):
        path = os.path.join(r"d:\NearER\portfolio_assets", img_filename)
        if os.path.exists(path):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run()
            run.add_picture(path, width=Inches(6.0))
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(10)
            r_c = p_cap.add_run(caption_text)
            r_c.font.size = Pt(9)
            r_c.font.italic = True
            r_c.font.color.rgb = GRAY

    # Section 1: Executive Profile
    add_h1("1. Executive Summary & Engineering Capabilities")
    add_body("I engineer robust, scalable, and visually stunning software products tailored for client business growth. From cross-platform mobile applications to multi-tenant web portals, real-time GPS tracking engines, and cloud backends, I deliver production-ready platforms with clean maintainable code.")

    # Section 2: ALL 6 EXACT PROJECTS SHOWCASE
    add_h1("2. Featured Projects & Architecture Diagrams")

    # 1. NearER
    add_h2("Project 1: NearER — Intelligent Emergency Healthcare System (🥈 2nd Prize Winner)")
    add_body("NearER connects patients, ambulances, hospitals, and control booths to eliminate emergency transit delays by routing cases to the 'Nearest Suitable and Available Hospital' based on live ER bed, ICU, ventilator, and specialist availability.")
    add_bullet("Key Stack:", "Flutter, Next.js 14, TypeScript, Node.js, WebSockets, Supabase / PostGIS.")
    add_image_visual("visual_nearer.png", "Figure 2.1: NearER Multi-Portal Ecosystem & Decision Engine Architecture")

    # 2. Zerviqo
    add_h2("Project 2: Zerviqo — High-Performance Service & Enterprise Automation Platform")
    add_body("Zerviqo is an end-to-end service management and dispatch ecosystem enabling automated task booking, real-time WebSocket job tracking, vendor partner onboarding, and executive KPI analytics.")
    add_bullet("Key Stack:", "Next.js 14, TypeScript, Node.js, WebSockets, PostgreSQL, Tailwind CSS.")
    add_image_visual("visual_zerviqo.png", "Figure 2.2: Zerviqo Service Dispatch & Enterprise Automation Architecture")

    # 3. Delivery App
    add_h2("Project 3: Delivery App — On-Demand GPS Tracking & Rider Dispatch System")
    add_body("A high-concurrency delivery logistics platform featuring live customer ordering, driver turn-by-turn navigation apps, dynamic traffic ETA calculations, and automated rider earnings settlement.")
    add_bullet("Key Stack:", "Flutter (Rider & Customer Apps), Go / Node.js, PostGIS geospatial engine, Socket.io, Mapbox GL.")
    add_image_visual("visual_delivery_app.png", "Figure 2.3: Delivery App Live Telemetry & Dynamic Dispatch Workflow")

    # 4. Ecomesh
    add_h2("Project 4: Ecomesh — Multi-Merchant E-Commerce Mesh Marketplace")
    add_body("Ecomesh is a multi-vendor e-commerce mesh network connecting distributed merchant storefronts, cross-warehouse inventory syncing, secure multi-party split payments, and automated order fulfillment.")
    add_bullet("Key Stack:", "React, Next.js, Node.js microservices, PostgreSQL, Stripe Connect, Redis cache.")
    add_image_visual("visual_ecomesh.png", "Figure 2.4: Ecomesh Multi-Merchant Network & Split Settlement Architecture")

    # 5. DMBH
    add_h2("Project 5: DMBH — Digital Medical & Business Healthcare Suite")
    add_body("DMBH is a comprehensive hospital & medical management portal unifying patient electronic records, clinical triage schedules, automated insurance claims billing, and executive health analytics.")
    add_bullet("Key Stack:", "Next.js, TypeScript, PostgreSQL, Prisma ORM, Tailwind CSS, Docker.")
    add_image_visual("visual_dmbh.png", "Figure 2.5: DMBH Digital Medical & Hospital Management Architecture")

    # 6. Highway Mirchi
    add_h2("Project 6: Highway Mirchi — Highway Dining, Online Ordering & Hospitality System")
    add_body("Highway Mirchi is a specialized food & hospitality platform enabling en-route highway meal pre-ordering, kitchen display system (KDS) queueing, express table POS ordering, and traveler loyalty rewards.")
    add_bullet("Key Stack:", "Flutter Mobile App, Next.js Kitchen Portal, Node.js WebSockets, Supabase, Tailwind CSS.")
    add_image_visual("visual_highway_mirchi.png", "Figure 2.6: Highway Mirchi En-Route Ordering & POS Hospitality Workflow")

    # Section 3: Services Offered
    add_h1("3. Client Service Offerings")
    tbl_srv = doc.add_table(rows=1, cols=3)
    tbl_srv.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_srv = tbl_srv.rows[0]
    hdr_srv.cells[0].text = "Service Area"
    hdr_srv.cells[1].text = "Deliverables"
    hdr_srv.cells[2].text = "Target Clients"
    style_table_header(hdr_srv, "0A2540")
    
    services_data = [
        ("Mobile App Development", "Cross-platform iOS & Android apps (Flutter/React Native) with offline sync, GPS & push notifications.", "Startups & mobile product teams"),
        ("Full-Stack Web Applications", "Next.js/React portals, SaaS platforms, e-commerce, custom dashboards with SEO & high performance.", "Companies launching web platforms"),
        ("Real-Time Systems & APIs", "WebSockets, live tracking, telemetry streaming, microservices, REST & GraphQL APIs.", "Logistics, healthcare & delivery apps"),
        ("UI/UX Design Systems", "Figma wireframing, high-fidelity light UI design, micro-animations, responsive design.", "Clients wanting premium look & feel"),
        ("Cloud & Database Architecture", "Supabase, PostgreSQL, PostGIS, Docker, AWS, security audits & database optimization.", "Scaling backend infrastructure")
    ]
    
    for row_idx, (cat, deliv, ideal) in enumerate(services_data):
        row = tbl_srv.add_row()
        row.cells[0].text = cat
        row.cells[1].text = deliv
        row.cells[2].text = ideal
        for i, cell in enumerate(row.cells):
            set_cell_background(cell, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            for r in p.runs:
                r.font.name = 'Calibri'
                r.font.size = Pt(9.5)
                if i == 0:
                    r.font.bold = True
                    r.font.color.rgb = NAVY

    # Section 4: Skills Matrix
    add_h1("4. Technical Skills Matrix")
    tbl_sk = doc.add_table(rows=1, cols=2)
    tbl_sk.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_sk = tbl_sk.rows[0]
    hdr_sk.cells[0].text = "Domain / Layer"
    hdr_sk.cells[1].text = "Technologies Mastered"
    style_table_header(hdr_sk, "0052FF")
    
    skills_data = [
        ("Languages", "TypeScript, JavaScript, Dart, Python, SQL, HTML5, CSS3"),
        ("Frontend / Web", "Next.js 14, React.js, Tailwind CSS, Framer Motion, Redux Toolkit, WebRTC"),
        ("Mobile Development", "Flutter, React Native, iOS & Android Native SDKs, Google Maps"),
        ("Backend & APIs", "Node.js, Express, Go, REST APIs, WebSockets, Socket.io, Microservices"),
        ("Databases & Storage", "PostgreSQL, Supabase, PostGIS, Redis, Prisma ORM, MongoDB"),
        ("DevOps & Tools", "Vercel, Docker, AWS, Firebase, Git/GitHub, Postman, Figma")
    ]
    
    for row_idx, (domain, tech) in enumerate(skills_data):
        row = tbl_sk.add_row()
        row.cells[0].text = domain
        row.cells[1].text = tech
        for i, cell in enumerate(row.cells):
            set_cell_background(cell, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            for r in p.runs:
                r.font.name = 'Calibri'
                r.font.size = Pt(9.5)
                if i == 0:
                    r.font.bold = True
                    r.font.color.rgb = NAVY

    # Section 5: Work With Me
    add_h1("5. Contact & Collaboration")
    add_callout_box(doc, [
        "Ready to build your web app, mobile app, or real-time platform?",
        "Email: your.email@example.com  |  Phone / WhatsApp: +91-XXXXX-XXXXX",
        "Live Portfolio URL: http://localhost:3000/portfolio  |  GitHub: github.com/your-username",
        "Available for contract roles, freelance client projects, and full-stack technical consulting."
    ], title="HIRE ME FOR YOUR NEXT PROJECT", border_color="0052FF", bg_color="F8FAFC")

    output_file = r"d:\NearER\NearER_Client_Portfolio.docx"
    doc.save(output_file)
    print(f"Updated user projects portfolio docx created successfully: {output_file}")

build_portfolio()
